import os
import time
import pandas as pd
import openpyxl
import traceback
import re
import pyotp
import selenium

from datetime import datetime
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


# =========================================================
# CONFIG
# =========================================================

EXCEL_FILE = r"C:\Users\mages\Desktop\Automation\HFA_Extarction\HFA_Extarction.xlsx"
DOWNLOAD_PATH = r"C:\Users\mages\Desktop\Automation\HFA_Extarction\Claim_Form"

URL = "https://pvpm.practicevelocity.com/"


# =========================================================
# EXCEL
# =========================================================

login_df = pd.read_excel(
    EXCEL_FILE,
    sheet_name="Login",
    dtype=str
).fillna("")

claimfrom_df = pd.read_excel(
    EXCEL_FILE,
    sheet_name="Claimfrom",
    dtype=str
).fillna("")


# =========================================================
# CHROME DRIVER
# =========================================================

def start_driver():
    os.makedirs(DOWNLOAD_PATH, exist_ok=True)

    options = Options()
    options.add_argument("--start-maximized")
    options.add_experimental_option("prefs", {
        "download.default_directory": DOWNLOAD_PATH,
        "download.prompt_for_download": False,
        "download.directory_upgrade": True,
        "plugins.always_open_pdf_externally": True,
    })

    driver = webdriver.Chrome(options=options)
    return driver


# =========================================================
# LOGIN
# =========================================================

def login(driver, username, password, otp):
    wait = WebDriverWait(driver, 30)

    driver.get(URL)

    wait.until(EC.presence_of_element_located((By.ID, "txtLogin"))).send_keys(username)
    wait.until(EC.element_to_be_clickable((By.ID, "btnNext"))).click()

    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='form20']/div[2]/input"))).click()

    wait.until(EC.presence_of_element_located((By.ID, "input54"))).send_keys(password)
    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='form46']/div[2]/input"))).click()

    wait.until(EC.presence_of_element_located((By.ID, "input80"))).send_keys(otp)
    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='form72']/div[2]/input"))).click()


def Patient(driver):
    wait = WebDriverWait(driver, 30)

    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='tdMenuBarItemPatient']/a"))).click()
    print("Patient clicked")

    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='menu_Patient_PatAR']"))).click()
    print("Receivable clicked")


def Select_Practice(driver, practice, payer_id, invoice):
    wait = WebDriverWait(driver, 30)

    wait.until(EC.presence_of_element_located((By.XPATH, "//*[@id='ddlPractice']"))).send_keys(practice)
    print(f"Practice Selected: {practice}")

    print("invoice", invoice)
    wait.until(EC.presence_of_element_located((By.XPATH, "//*[@id='txtInvNum']"))).send_keys(invoice, Keys.ENTER)

    if driver.find_elements(By.XPATH, "//*[@id='ShowCMSLine']/label/span"):
        driver.find_element(By.XPATH, "//*[@id='ShowCMSLine']/label/span").click()
        print("Data Stamp unchecked")

        wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='ShowCMSLine']/input"))).click()
        print("CMS500 Clicked")
        wait.until(lambda d: len(d.window_handles) > 1)
        driver.switch_to.window(driver.window_handles[1])
        wait.until(lambda d: d.execute_script("return document.readyState") == "complete")
        return True

    else:
        wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='btnBack']"))).click()
        print("CMS500 not found")
        return False


def download_claim_form(driver, practice, payer_id, invoice, claim_from):
    wait = WebDriverWait(driver, 30)

    practice_path = os.path.join(DOWNLOAD_PATH, practice)
    os.makedirs(practice_path, exist_ok=True)

    # Send downloads to the practice folder
    driver.execute_cdp_cmd(
        "Browser.setDownloadBehavior",
        {"behavior": "allow", "downloadPath": practice_path}
    )

    before = set(os.listdir(practice_path))

    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='save']"))).click()
    print("Claim Form Downloaded")

    # Wait until a new PDF appears and no .crdownload is left
    WebDriverWait(driver, 60).until(
        lambda d: not any(f.endswith(".crdownload") for f in os.listdir(practice_path))
        and any(f.lower().endswith(".pdf") for f in set(os.listdir(practice_path)) - before)
    )
    
    driver.switch_to.window(driver.window_handles[0])
    
    wait.until(EC.element_to_be_clickable((By.XPATH, "//*[@id='btnBack']"))).click()

    old_file = max(
        [
            os.path.join(practice_path, f)
            for f in set(os.listdir(practice_path)) - before
            if f.lower().endswith(".pdf")
        ],
        key=os.path.getmtime
    )

    new_file_path = os.path.join(
        practice_path,
        f"{payer_id}_{invoice}_{claim_from}.pdf"
    )

    if os.path.exists(new_file_path):
        os.remove(new_file_path)

    os.rename(old_file, new_file_path)


def update_status(index, value):
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Claimfrom"]
    ws.cell(row=index + 2, column=4).value = value
    wb.save(EXCEL_FILE)


# =========================================================
# MAIN
# =========================================================

def main():
    # ---- LOGIN SHEET READ ----
    username = password = security_key = ""

    for index, row in login_df.iterrows():

        username = row["Username"].strip()
        password = row["Password"].strip()
        security_key = row["Security Key"].strip().replace(" ", "").upper()

        if not username or not security_key:
            print(f"Row {index + 2}: Username or Security Key missing")
            continue

        try:
            otp = pyotp.TOTP(security_key).now()
            break
        except Exception as e:
            print(f"Row {index + 2}: Invalid Security Key - {e}")

    else:
        print("No valid login found")
        return

    driver = start_driver()

    try:
        login(driver, username, password, otp)
        Patient(driver)

        # ---- PRACTICE SHEET READ ----
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb["Claimfrom"]

        for row_index in range(2, ws.max_row + 1):

            practice = str(ws.cell(row=row_index, column=1).value or "").strip()
            payer_id = str(ws.cell(row=row_index, column=2).value or "").strip()
            invoice = str(ws.cell(row=row_index, column=3).value or "").strip()
            claim_form_status = str(ws.cell(row=row_index, column=4).value or "").strip()
            claim_from = str(ws.cell(row=row_index, column=5).value or "").strip()

            if not practice or not invoice:
                continue

            if claim_form_status.lower() == "downloaded":
                continue

            try:
                if Select_Practice(driver, practice, payer_id, invoice):

                    download_claim_form(
                        driver,
                        practice,
                        payer_id,
                        invoice,
                        claim_from
                    )

                    ws.cell(row=row_index, column=4).value = "Downloaded"
                    wb.save(EXCEL_FILE)

                    print(f"Row {row_index}: Downloaded")

                else:
                    ws.cell(row=row_index, column=4).value = "CMS500 Not Found"
                    wb.save(EXCEL_FILE)

                    print(f"Row {row_index}: CMS500 Not Found")

            except Exception as e:

                print(f"Row {row_index}: Download Failed - {e}")
                traceback.print_exc()

                ws.cell(row=row_index, column=4).value = "Failed"
                wb.save(EXCEL_FILE)

                # return to main tab if popup was left open
                while len(driver.window_handles) > 1:
                    driver.switch_to.window(driver.window_handles[-1])
                    driver.close()

                driver.switch_to.window(driver.window_handles[0])

    finally:
        driver.quit()

if __name__ == "__main__":
    main()