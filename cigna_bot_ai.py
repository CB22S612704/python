"""
Cigna for HCP - ADVANCED AI-assisted claim automation + analytics

Install:   pip install selenium openpyxl anthropic
API key:   setx ANTHROPIC_API_KEY "sk-ant-..."   (Windows; reopen terminal)
Run:       python cigna_bot_ai.py [--headless] [--no-ai]

What is new vs cigna_bot.py
  1. AI status writer   - remark-code explanation -> short billing-note phrase (when not in Master)
  2. AI claim analysis  - category, root cause, next action, priority per claim
  3. Rule-based checks  - underpayment / variance vs billed amount (works even with --no-ai)
  4. Analysis sheet     - totals, status mix, top denial codes, AI executive summary
  5. Reliability        - retries, fallback selectors, screenshots on error, log file, resume
  6. Privacy            - ONLY de-identified data (CPT, remark codes, amounts) is sent to the AI.
                          Names, DOB, IDs, claim numbers never leave your machine.
                          (Still confirm your organisation's HIPAA / BAA policy before enabling AI.)
"""
import os, re, sys, json, time, pickle, hashlib, logging, functools
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException

# ============================================================ CONFIG
EXCEL_PATH = r"C:\Users\Automation22\Desktop\Automation\Cigna\Cigna.xlsx"
WORK_DIR = r"C:\Users\Automation22\Desktop\Automation\Cigna"
SESSION_DIR = os.path.join(WORK_DIR, "session")
SHOT_DIR = os.path.join(WORK_DIR, "screenshots")
COOKIE_FILE = os.path.join(SESSION_DIR, "cigna_cookies.pkl")
AI_CACHE_FILE = os.path.join(WORK_DIR, "ai_cache.json")
LOGIN_URL = "https://cignaforhcp.cigna.com/app/login"
AI_MODEL = "claude-sonnet-5-5"
WAIT = 20
USE_AI = "--no-ai" not in sys.argv
HEADLESS = "--headless" in sys.argv

# Input sheet columns (1-based) - first 13 are your original spec
(C_INV, C_NAME, C_DOB, C_INS, C_ID, C_DOS, C_CPT_IN, C_BILLED, C_NOTES1, C_NOTES2,
 C_CPT_OUT, C_REMARK, C_STATUS) = range(1, 14)
# extra output columns added by this version
C_CATEGORY, C_ACTION, C_PRIORITY, C_VARIANCE, C_SEARCH, C_CONF = range(14, 20)
EXTRA_HEADERS = {C_CATEGORY: "AI Category", C_ACTION: "Next Action", C_PRIORITY: "Priority",
                 C_VARIANCE: "Billed - Paid", C_SEARCH: "Search Type", C_CONF: "Match Confidence"}

os.makedirs(SESSION_DIR, exist_ok=True); os.makedirs(SHOT_DIR, exist_ok=True)
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                    handlers=[logging.FileHandler(os.path.join(WORK_DIR, "cigna_bot.log"), encoding="utf-8"),
                              logging.StreamHandler()])
log = logging.getLogger("cigna")

# ============================================================ XPATHS (from your spec)
BASE = "/html/body/cigna-root/cigna-layout-wrapper/cigna-layout-secure/div/main/div/cigna-page"
X = {
    "cookie_close": '//*[@id="onetrust-close-btn-container"]/button',
    "login_btn": "/html/body/cigna-root/cigna-layout-wrapper/cigna-layout-public/div/main/cigna-page/div/div[1]/div/cigna-login-form-auth0/div/div[2]/button",
    "username": '//*[@id="username-label"]',
    "next": "/html/body/main/section/div/div/div/div[1]/div/form/div[4]/button",
    "password": '//*[@id="password"]',
    "continue": "/html/body/main/section/div/div/div/form/div[5]/button",
    "otp": '//*[@id="code"]',
    "remember": "/html/body/main/section/div/div/div/div/div/form/div[3]/label/span",
    "otp_continue": "/html/body/main/section/div/div/div/div/div/form/div[4]/button",
    "claims": "/html/body/cigna-root/cigna-layout-wrapper/cigna-layout-secure/div/header/cigna-primary-nav/nav/div[3]/div/cigna-nav-item/a",
    "search_claim": "/html/body/cigna-root/cigna-layout-wrapper/cigna-layout-secure/div/header/cigna-primary-nav/nav/div[3]/cigna-nav-links/div[1]/cigna-nav-item/a",
    "t1_radio": '//*[@id="results-top"]/div[1]/cigna-search-form/div/cigna-tab-group/div/cigna-search-type-patient/form/div[1]/fieldset/div[1]/div[1]/label/i',
    "t2_radio": '//*[@id="results-top"]/div[1]/cigna-search-form/div/cigna-tab-group/div/cigna-search-type-patient/form/div[1]/fieldset/div[1]/div[2]/label/i',
    "dob": '//*[@id="dob"]', "member_id": '//*[@id="memberId"]',
    "first": '//*[@id="firstName"]', "last": '//*[@id="lastName"]',
    "search": '//*[@id="results-top"]/div[1]/cigna-search-form/div/cigna-tab-group/div/cigna-search-type-patient/form/div[2]/fieldset/div/div[6]/button[1]',
    "res_rows": BASE + "/cigna-search-wrapper/cigna-main-screen/cigna-results-table/table/tbody/tr",
    "return": BASE + "/cigna-claim-details-screen/cigna-breadcrumbs/div/div[1]/a",
}
D = BASE + "/cigna-claim-details-screen/cigna-claims-details/div/cigna-claims-details-medical"
INFO = D + "/cigna-medical-claims-information/div/cigna-tab-group/div/div/div"
X.update({
    "claim_no": INFO + "/div[1]/table/tbody/tr[1]/td",
    "date_processed": INFO + "/div[1]/table/tbody/tr[6]/td",
    "status": INFO + "/div[1]/table/tbody/tr[7]/td/cigna-claim-status/cigna-status-badge/div/div[2]",
    "amt_due": INFO + "/div[2]/table/tbody/tr[1]/td",
    "amt_paid": INFO + "/div[2]/table/tbody/tr[2]/td",
    "provider_pay": INFO + "/div[2]/table/tbody/tr[3]/td",
    "patient_resp": INFO + "/div[2]/table/tbody/tr[4]/td",
    "proc_rows": D + "/div[4]/cigna-medical-procedure-details/div/table/tbody/tr",
    "expl_blocks": D + "/div[5]/cigna-medical-explanation-details/div[1]/div",
})
# Fallback (label based) selectors in case Cigna shifts the DOM - tried when the main XPath fails
FALLBACK = {
    "claim_no": ["//th[contains(.,'Claim/Reference')]/following-sibling::td[1]"],
    "date_processed": ["//th[contains(.,'Date Processed')]/following-sibling::td[1]"],
    "amt_paid": ["//th[contains(.,'Claim amount paid')]/following-sibling::td[1]"],
    "provider_pay": ["//th[contains(.,'Total provider payment')]/following-sibling::td[1]"],
    "amt_due": ["//th[contains(.,'Claim amount due')]/following-sibling::td[1]"],
    "patient_resp": ["//th[contains(.,'Patient responsibility')]/following-sibling::td[1]"],
}

# ============================================================ UTILS
def clean(s): return re.sub(r"\s+", " ", str(s or "")).strip()

def money(s):
    try: return float(re.sub(r"[^\d.\-]", "", str(s or "")) or 0)
    except ValueError: return 0.0

def norm_date(v):
    if isinstance(v, datetime): return v.strftime("%m/%d/%Y")
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", str(v or ""))
    if not m: return ""
    mo, d, y = m.groups(); y = "20" + y if len(y) == 2 else y
    return f"{int(mo):02d}/{int(d):02d}/{y}"

def retry(times=3, delay=2):
    def deco(fn):
        @functools.wraps(fn)
        def wrap(*a, **k):
            for i in range(times):
                try: return fn(*a, **k)
                except Exception as e:
                    if i == times - 1: raise
                    log.warning("retry %s (%s): %s", fn.__name__, i + 1, e); time.sleep(delay)
        return wrap
    return deco

# ============================================================ AI LAYER
class AI:
    """Thin Claude wrapper: JSON in/out, disk cache, silent fallback if no key / error."""
    def __init__(self, enabled):
        self.client = None
        self.cache = json.load(open(AI_CACHE_FILE)) if os.path.exists(AI_CACHE_FILE) else {}
        if enabled and os.environ.get("ANTHROPIC_API_KEY"):
            try:
                import anthropic; self.client = anthropic.Anthropic()
                log.info("AI enabled (%s)", AI_MODEL)
            except ImportError:
                log.warning("pip install anthropic to enable AI")
        elif enabled:
            log.warning("ANTHROPIC_API_KEY not set - running rule-based only")

    @property
    def on(self): return self.client is not None

    def ask_json(self, system, payload, max_tokens=400):
        if not self.on: return None
        key = hashlib.sha256((system + json.dumps(payload, sort_keys=True)).encode()).hexdigest()
        if key in self.cache: return self.cache[key]
        try:
            r = self.client.messages.create(
                model=AI_MODEL, max_tokens=max_tokens, system=system + "\nReply with ONLY valid JSON.",
                messages=[{"role": "user", "content": json.dumps(payload)}])
            txt = re.sub(r"^```(?:json)?|```$", "", r.content[0].text.strip(), flags=re.M).strip()
            out = json.loads(txt)
            self.cache[key] = out
            json.dump(self.cache, open(AI_CACHE_FILE, "w"))
            return out
        except Exception as e:
            log.warning("AI call failed: %s", e); return None

    def short_status(self, cpt, code, explanation):
        r = self.ask_json(
            "You write medical-billing account notes. Condense the payer remark-code explanation into a "
            "terse note phrase (max 10 words, standard billing abbreviations, no PHI). "
            'Format: {"status": "..."}', {"cpt": cpt, "code": code, "explanation": explanation}, 120)
        return clean(r.get("status")) if r else None

    def analyze_claim(self, facts):
        return self.ask_json(
            "You are a senior US medical-billing analyst. Given de-identified claim facts, classify it. "
            'Return {"category": one of ["Paid in full","Partially paid","Denied","Patient responsibility",'
            '"Pending/Other"], "root_cause": "<=15 words", "action": one of ["No action","Post payment",'
            '"Bill patient","Correct & resubmit","File appeal","Request records/auth","Call payer"], '
            '"priority": one of ["High","Medium","Low"]}', facts, 250)

    def executive_summary(self, stats):
        r = self.ask_json(
            "You are a revenue-cycle analyst. Write a 4-6 sentence executive summary of this batch of "
            'claim results: key findings, biggest denial drivers, recommended focus. {"summary": "..."}',
            stats, 500)
        return r.get("summary") if r else None

# ============================================================ RULE-BASED ANALYSIS (fallback + variance)
def rule_analysis(billed, paid, resp, denied):
    if denied: return dict(category="Denied", action="File appeal", priority="High",
                           root_cause="All amounts zero on claim")
    if paid > 0 and billed and paid + resp + 0.01 < billed * 0.5:
        return dict(category="Partially paid", action="Call payer", priority="Medium",
                    root_cause="Payment well below billed amount")
    if resp > 0 and paid == 0: return dict(category="Patient responsibility", action="Bill patient",
                                           priority="Low", root_cause="Payer assigned full amount to patient")
    return dict(category="Paid in full" if paid else "Pending/Other",
                action="Post payment" if paid else "Call payer", priority="Low", root_cause="")

# ============================================================ MASTER SHEET
class Master:
    def __init__(self, ws):
        self.ws = ws; self.map = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0]: self.map[(clean(r[0]), clean(r[1]))] = clean(r[2])
    def get(self, cpt, code): return self.map.get((cpt, str(code)))
    def add(self, cpt, code, status):
        if (cpt, str(code)) in self.map: return                       # exists -> no duplicate
        row = 2
        while self.ws.cell(row, 1).value not in (None, ""): row += 1   # first empty row
        self.ws.cell(row, 1, cpt); self.ws.cell(row, 2, str(code)); self.ws.cell(row, 3, status)
        self.map[(cpt, str(code))] = status

# ============================================================ BROWSER
class Bot:
    def __init__(self):
        o = webdriver.ChromeOptions()
        o.add_argument(f"--user-data-dir={os.path.join(SESSION_DIR, 'profile')}")
        o.add_argument("--start-maximized")
        if HEADLESS: o.add_argument("--headless=new")
        self.d = webdriver.Chrome(options=o)

    def find(self, xp, t=WAIT): return WebDriverWait(self.d, t).until(EC.presence_of_element_located((By.XPATH, xp)))
    def click(self, xp, t=WAIT):
        el = WebDriverWait(self.d, t).until(EC.element_to_be_clickable((By.XPATH, xp)))
        try: el.click()
        except Exception: self.d.execute_script("arguments[0].click();", el)
    def type(self, xp, v): el = self.find(xp); el.clear(); el.send_keys(v)
    def count(self, xp): return len(self.d.find_elements(By.XPATH, xp))
    def text(self, key_or_xp, t=3):
        xps = [X.get(key_or_xp, key_or_xp)] + FALLBACK.get(key_or_xp, [])
        for xp in xps:
            try: return clean(self.find(xp, t).text)
            except TimeoutException: continue
        return ""
    def shot(self, name):
        p = os.path.join(SHOT_DIR, f"{datetime.now():%Y%m%d_%H%M%S}_{name}.png")
        try: self.d.save_screenshot(p)
        except Exception: pass
        return p
    def close_banner(self):
        try: self.click(X["cookie_close"], 3)
        except TimeoutException: pass

    # ---- login
    def save_cookies(self): pickle.dump(self.d.get_cookies(), open(COOKIE_FILE, "wb"))
    def load_cookies(self):
        if os.path.exists(COOKIE_FILE):
            for c in pickle.load(open(COOKIE_FILE, "rb")):
                try: self.d.add_cookie(c)
                except Exception: pass
    def logged_in(self, t=8):
        try: self.find(X["claims"], t); return True
        except TimeoutException: return False

    def login(self, user, pwd):
        user, pwd = re.sub(r"\s+", "", str(user)), re.sub(r"\s+", "", str(pwd))   # no spaces
        self.d.get(LOGIN_URL); self.load_cookies(); self.d.get(LOGIN_URL); self.close_banner()
        if self.logged_in(4): return log.info("Already logged in")
        try:
            self.click(X["login_btn"], 8)                                           # Method 2
            if self.logged_in(): self.save_cookies(); return log.info("Logged in via saved session")
        except TimeoutException: pass
        log.info("Full login (Method 1)")                                           # Method 1
        self.type(X["username"], user); self.click(X["next"])
        self.type(X["password"], pwd); self.click(X["continue"])
        self.type(X["otp"], input("Enter Cigna authentication code: ").strip())
        try: self.click(X["remember"], 5)
        except TimeoutException: pass
        self.click(X["otp_continue"])
        if not self.logged_in(25): raise RuntimeError("Login failed")
        self.save_cookies(); log.info("Login OK - session saved (~30 days)")

    # ---- search
    @retry(3)
    def goto_search(self): self.close_banner(); self.click(X["claims"]); self.click(X["search_claim"])
    def has_results(self):
        try: WebDriverWait(self.d, 10).until(lambda d: self.count(X["res_rows"]) > 0); return True
        except TimeoutException: return False

    def search_patient(self, first, last, dob, pid):
        self.goto_search()                                             # Type 1: DOB + ID
        self.click(X["t1_radio"]); self.type(X["dob"], dob); self.type(X["member_id"], str(pid)); self.click(X["search"])
        if self.has_results(): return "Type1"
        self.goto_search()                                             # Type 2: Name + DOB
        self.click(X["t2_radio"]); self.type(X["first"], first); self.type(X["last"], last)
        self.type(X["dob"], dob); self.click(X["search"])
        return "Type2" if self.has_results() else None

    # ---- claim details
    @retry(2)
    def read_claim(self):
        c = {k: self.text(k) for k in ("claim_no", "date_processed", "status", "amt_due", "amt_paid",
                                       "provider_pay", "patient_resp")}
        expl = {}
        for i in range(1, self.count(X["expl_blocks"]) + 1):
            b = f"{X['expl_blocks']}[{i}]"
            head, body = self.text(b + "/div[1]", 2), self.text(b + "/div[2]/p", 2)
            m = re.match(r"\s*(\w+)", head)
            if m: expl[m.group(1)] = body or head
        procs = []
        for i in range(1, self.count(X["proc_rows"]) + 1):
            r = f"{X['proc_rows']}[{i}]"
            procs.append(dict(cpt=self.text(r + "/td[1]", 2), dos=self.text(r + "/td[2]", 2),
                              remarks=[x for x in re.split(r"[,\s]+", self.text(r + "/td[12]", 2)) if x]))
        c["procs"], c["expl"] = procs, expl
        return c

# ============================================================ NOTES BUILDER
def build_notes(claim, cpt_filter, excel_dos, master, ai):
    procs = [p for p in claim["procs"] if not cpt_filter or p["cpt"] == cpt_filter]
    if not procs: return None
    due, paid, prov, resp = (money(claim[k]) for k in ("amt_due", "amt_paid", "provider_pay", "patient_resp"))
    denied = (due == paid == prov == resp == 0)
    word = "DND" if denied else "PAID"
    amt = prov or paid
    dos = excel_dos or procs[0]["dos"]
    head = f"{dos} CLM PD ${amt:,.2f} THRU CHK #{claim['claim_no']} ISSUED ON {claim['date_processed']}."
    parts, codes = [], []
    for p in procs:
        for code in (p["remarks"] or [""]):
            expl = claim["expl"].get(code, "")
            status = master.get(p["cpt"], code)
            if status is None:                                         # not in Master -> AI/plain text, then add
                status = (ai.short_status(p["cpt"], code, expl) if ai.on and expl else None) or expl
                if code: master.add(p["cpt"], code, status)
            parts.append(clean(f"{p['cpt']} {word} {code} - {status}").rstrip(" -"))
            codes.append(code)
    return dict(note=head + " " + "; ".join(parts), cpt=",".join(p["cpt"] for p in procs),
                remark=",".join(dict.fromkeys(codes)), status=claim["status"],
                paid=amt, resp=resp, denied=denied, procs=procs, claim=claim)

# ============================================================ PER-ROW PIPELINE
def analyze(res, billed, ai):
    facts = dict(cpts=[p["cpt"] for p in res["procs"]], claim_status=res["status"], billed=billed,
                 paid=res["paid"], patient_resp=res["resp"], denied=res["denied"],
                 remarks=[{"code": c, "text": res["claim"]["expl"].get(c, "")}
                          for p in res["procs"] for c in p["remarks"]])
    return (ai.analyze_claim(facts) if ai.on else None) or rule_analysis(billed, res["paid"], res["resp"], res["denied"])

def process_row(bot, ws, r, master, ai):
    g = lambda c: ws.cell(r, c).value
    name = str(g(C_NAME) or "")
    first, last = ([clean(x) for x in name.split(",", 1)] if "," in name
                   else (name.split(" ", 1) + [""])[:2])
    dob, pid = norm_date(g(C_DOB)), g(C_ID)
    xdos, xcpt, billed = norm_date(g(C_DOS)), clean(g(C_CPT_IN)), money(g(C_BILLED))

    if not (dob and pid) and not (first and last and dob):
        ws.cell(r, C_STATUS, "Invalid input (missing DOB/ID/name)"); return
    stype = bot.search_patient(first, last, dob, pid)
    if not stype: ws.cell(r, C_STATUS, "Patient Not Found"); return
    ws.cell(r, C_SEARCH, stype)

    rows = [(norm_date(bot.text(f"{X['res_rows']}[{i}]/td[6]", 2)), i)
            for i in range(1, bot.count(X["res_rows"]) + 1)]

    def open_read(i):
        bot.click(f"{X['res_rows']}[{i}]/td[2]/div/a"); c = bot.read_claim()
        bot.click(X["return"]); bot.has_results(); return c

    chosen, target = None, None
    if xdos and xcpt:                                                  # Notes1: DOS match + CPT match
        for d, i in rows:
            if d == xdos and (res := build_notes(open_read(i), xcpt, xdos, master, ai)):
                chosen, target = res, C_NOTES1; break
    if chosen: ws.cell(r, C_CONF, "High (DOS+CPT)")
    else:                                                              # Notes2: last two claims
        for d, i in rows[-2:]:
            if res := build_notes(open_read(i), None, xdos, master, ai):
                chosen, target = res, C_NOTES2
        ws.cell(r, C_CONF, "Low (fallback claims)") if chosen else None
    if not chosen: ws.cell(r, C_STATUS, "No Matching Claim"); return

    a = analyze(chosen, billed, ai)
    ws.cell(r, target, chosen["note"]); ws.cell(r, C_CPT_OUT, chosen["cpt"])
    ws.cell(r, C_REMARK, chosen["remark"]); ws.cell(r, C_STATUS, chosen["status"])
    ws.cell(r, C_CATEGORY, a.get("category")); ws.cell(r, C_ACTION, a.get("action"))
    ws.cell(r, C_PRIORITY, a.get("priority")); ws.cell(r, C_VARIANCE, round(billed - chosen["paid"], 2))

# ============================================================ ANALYSIS SHEET
def write_analysis(wb, inp, ai):
    rows = [[c.value for c in row] for row in inp.iter_rows(min_row=2, max_col=C_CONF) if row[0].value]
    cat = Counter(r[C_CATEGORY - 1] or "Not processed" for r in rows)
    act = Counter(r[C_ACTION - 1] for r in rows if r[C_ACTION - 1])
    codes = Counter(c for r in rows for c in str(r[C_REMARK - 1] or "").split(",") if c)
    unf = Counter(r[C_STATUS - 1] for r in rows if str(r[C_STATUS - 1] or "").startswith(("Patient Not", "No Match", "Invalid", "Error")))
    billed = sum(money(r[C_BILLED - 1]) for r in rows); var = sum(money(r[C_VARIANCE - 1]) for r in rows)
    stats = dict(total=len(rows), categories=dict(cat), actions=dict(act), top_remark_codes=codes.most_common(8),
                 unresolved=dict(unf), total_billed=round(billed, 2), total_variance=round(var, 2))
    name = "Analysis"
    if name in wb.sheetnames: del wb[name]
    ws = wb.create_sheet(name)
    ws.append(["Cigna batch analysis", f"{datetime.now():%Y-%m-%d %H:%M}"]); ws.append([])
    for title, data in (("Category", cat), ("Next action", act), ("Top remark codes", dict(codes.most_common(8))),
                        ("Unresolved", unf)):
        ws.append([title, "Count"]); [ws.append([k, v]) for k, v in data.items()]; ws.append([])
    ws.append(["Total billed", billed]); ws.append(["Total billed - paid", var]); ws.append([])
    summary = ai.executive_summary(stats) if ai.on else None
    ws.append(["AI executive summary", summary or "(AI off - see counts above)"])
    ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 90

# ============================================================ MAIN
def main():
    wb = load_workbook(EXCEL_PATH)
    inp, master = wb["Input"], Master(wb["Master"])
    for c, h in EXTRA_HEADERS.items():
        if not inp.cell(1, c).value: inp.cell(1, c, h)
    user, pwd = wb["Login"].cell(2, 1).value, wb["Login"].cell(2, 2).value
    ai, bot = AI(USE_AI), Bot()
    try:
        bot.login(user, pwd)
        for r in range(2, inp.max_row + 1):
            if not inp.cell(r, C_INV).value or inp.cell(r, C_STATUS).value: continue   # resume-safe
            log.info("Row %s | Invoice %s", r, inp.cell(r, C_INV).value)
            try: process_row(bot, inp, r, master, ai)
            except Exception as e:
                log.exception("Row %s failed", r); bot.shot(f"row{r}")
                inp.cell(r, C_STATUS, f"Error: {str(e)[:80]}")
                try: bot.goto_search()
                except Exception: pass
            wb.save(EXCEL_PATH)
        write_analysis(wb, inp, ai)
    finally:
        try: bot.save_cookies()
        except Exception: pass
        wb.save(EXCEL_PATH); bot.d.quit()
    log.info("Done")

if __name__ == "__main__":
    main()
